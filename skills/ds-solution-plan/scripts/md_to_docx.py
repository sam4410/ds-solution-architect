#!/usr/bin/env python3
"""Convert a DS Solution Plan written in Markdown into a styled Word (.docx) document.

Usage:
    python md_to_docx.py <input.md> [output.docx]

Supported Markdown subset (matches solution-plan-template.md):
  - Optional metadata block at the top between '---' lines (key: value) -> cover page
  - #, ##, ###, #### headings
  - Paragraphs with **bold**, *italic*, `code`, [text](url)
  - Bullets (-, *, +) with nesting by indentation, numbered lists (1.)
  - Pipe tables (header + separator row)
  - Blockquotes (> ...) rendered as shaded call-outs
  - Fenced code blocks (```) rendered in monospace (e.g. ASCII timelines)
  - Horizontal rules (---) ignored outside the metadata block

Requires: python-docx  (pip install python-docx)
"""
from __future__ import annotations

import datetime as _dt
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.stderr.write(
        "python-docx is not installed. Install it with:\n"
        "    pip install python-docx\n"
        "(or: pip install --user python-docx / pip install --break-system-packages python-docx)\n"
    )
    sys.exit(2)

# ---------------------------------------------------------------- styling
PRIMARY = RGBColor(0x1F, 0x3A, 0x5F)     # deep navy
ACCENT = RGBColor(0x2E, 0x75, 0xB6)      # blue
MUTED = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "1F3A5F"
ZEBRA_FILL = "F2F5F9"
CALLOUT_FILLS = {
    "assumption": "FFF4E5",
    "risk": "FDECEA",
    "open question": "EAF2FB",
    "note": "F2F2F2",
    "decision": "EAF7EE",
}
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

INLINE_RE = re.compile(
    r"(\*\*[^*]+?\*\*|__[^_]+?__|\*[^*\s][^*]*?\*|`[^`]+?`|\[[^\]]+\]\([^)]+\))"
)
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBER_RE = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")  # number re-extracted in convert()
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
STATUS_RE = re.compile(r"\[(Confirmed|Assumed|To investigate|TBD[^\]]*)\]", re.I)


# ---------------------------------------------------------------- helpers
def _shade(element, fill: str) -> None:
    """Apply background shading to a cell (tcPr) or paragraph (pPr)."""
    pr = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _left_border(paragraph, color: str = "2E75B6") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    bdr.append(left)
    pPr.append(bdr)


def _add_field(paragraph, instr: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = placeholder
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr_el, fld_sep, txt, fld_end):
        run._r.append(el)


def _repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _set_update_fields_on_open(doc) -> None:
    settings = doc.settings.element
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.append(el)


def add_inline(paragraph, text: str, size: Pt | None = None, color: RGBColor | None = None, bold_all=False) -> None:
    """Render inline markdown into runs."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            _run(paragraph, text[pos:m.start()], size=size, color=color, bold=bold_all)
        tok = m.group(0)
        if tok.startswith("**") or tok.startswith("__"):
            _run(paragraph, tok[2:-2], bold=True, size=size, color=color)
        elif tok.startswith("`"):
            _run(paragraph, tok[1:-1], mono=True, size=size, color=color, bold=bold_all)
        elif tok.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok).groups()
            _run(paragraph, label, size=size, color=ACCENT, bold=bold_all, underline=True)
            if not url.startswith("#"):
                _run(paragraph, f" ({url})", size=size, color=MUTED)
        else:
            _run(paragraph, tok[1:-1], italic=True, size=size, color=color, bold=bold_all)
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], size=size, color=color, bold=bold_all)


def _run(paragraph, text, bold=False, italic=False, mono=False, size=None, color=None, underline=False):
    # Highlight status tags such as [Confirmed] / [Assumed] / [To investigate]
    parts = STATUS_RE.split(text) if STATUS_RE.search(text) else [text]
    if len(parts) > 1:
        for i, part in enumerate(parts):
            if not part:
                continue
            if i % 2 == 1:
                tag = part.lower()
                col = (RGBColor(0x1E, 0x7B, 0x34) if tag.startswith("confirmed")
                       else RGBColor(0xB3, 0x6B, 0x00) if tag.startswith("assumed")
                       else RGBColor(0xC0, 0x39, 0x2B))
                r = paragraph.add_run(f"[{part}]")
                r.bold = True
                r.font.color.rgb = col
                if size:
                    r.font.size = size
            else:
                _run(paragraph, part, bold, italic, mono, size, color, underline)
        return
    r = paragraph.add_run(text)
    r.bold = bold or None
    r.italic = italic or None
    r.underline = underline or None
    if mono:
        r.font.name = MONO_FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
    if size:
        r.font.size = size
    if color is not None:
        r.font.color.rgb = color


# ---------------------------------------------------------------- document setup
def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)  # A4 portrait
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.1

    for lvl, size, color in ((1, 18, PRIMARY), (2, 15, PRIMARY), (3, 12.5, ACCENT), (4, 11, ACCENT)):
        st = styles[f"Heading {lvl}"]
        st.font.name = BODY_FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
    return doc


def add_cover(doc, meta: dict) -> None:
    title = meta.get("title", "Data Science Solution Plan")
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("DATA SCIENCE SOLUTION PLAN")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = ACCENT
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = PRIMARY
    p.paragraph_format.space_after = Pt(18)
    _left_border(p, HEADER_FILL)

    rows = [
        ("Client", meta.get("client")),
        ("Industry", meta.get("industry")),
        ("Prepared by", meta.get("prepared_by")),
        ("Date", meta.get("date") or _dt.date.today().isoformat()),
        ("Version", meta.get("version")),
        ("Plan depth", meta.get("depth")),
        ("Classification", meta.get("classification", "Confidential")),
    ]
    rows = [(k, v) for k, v in rows if v]
    tbl = doc.add_table(rows=len(rows), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (k, v) in enumerate(rows):
        c0, c1 = tbl.rows[i].cells
        c0.width, c1.width = Cm(4), Cm(11)
        c0.text = ""
        c1.text = ""
        rr = c0.paragraphs[0].add_run(k)
        rr.bold = True
        rr.font.color.rgb = MUTED
        add_inline(c1.paragraphs[0], str(v))
    doc.add_paragraph()
    note = doc.add_paragraph()
    add_inline(note, "*Status tags used throughout:* **[Confirmed]** stated by client or verified in data · "
                     "**[Assumed]** working assumption to validate · **[To investigate]** existence or usability unknown.",
               size=Pt(9), color=MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    h = doc.add_paragraph()
    hr = h.add_run("Contents")
    hr.bold = True
    hr.font.size = Pt(15)
    hr.font.color.rgb = PRIMARY
    toc = doc.add_paragraph()
    _add_field(toc, 'TOC \\o "1-3" \\h \\z \\u',
               "Right-click here and choose 'Update Field' to build the table of contents.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_header_footer(doc, meta: dict) -> None:
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_inline(hp, meta.get("title", "Data Science Solution Plan"), size=Pt(8), color=MUTED)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"{meta.get('classification', 'Confidential')}  ·  Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    _add_field(fp, "PAGE", "1")
    r = fp.add_run(" of ")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    _add_field(fp, "NUMPAGES", "1")


# ---------------------------------------------------------------- block renderers
def split_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    cells, buf, in_code = [], "", False
    i = 0
    while i < len(s):
        ch = s[i]
        if ch == "`":
            in_code = not in_code
        if ch == "\\" and i + 1 < len(s) and s[i + 1] == "|":
            buf += "|"
            i += 2
            continue
        if ch == "|" and not in_code:
            cells.append(buf.strip())
            buf = ""
        else:
            buf += ch
        i += 1
    cells.append(buf.strip())
    return cells


def add_table(doc, header: list[str], rows: list[list[str]]) -> None:
    ncols = len(header)
    font = Pt(9.5 if ncols <= 4 else 8.5 if ncols <= 6 else 7.5)
    tbl = doc.add_table(rows=1, cols=ncols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0]
    _repeat_header(hdr)
    for j, text in enumerate(header):
        cell = hdr.cells[j]
        cell.text = ""
        _shade(cell._tc, HEADER_FILL)
        add_inline(cell.paragraphs[0], text, size=font, color=RGBColor(0xFF, 0xFF, 0xFF), bold_all=True)
    for i, row in enumerate(rows):
        cells = tbl.add_row().cells
        row = (row + [""] * ncols)[:ncols]
        for j, text in enumerate(row):
            cells[j].text = ""
            if i % 2 == 1:
                _shade(cells[j]._tc, ZEBRA_FILL)
            parts = re.split(r"<br\s*/?>", text)
            for k, part in enumerate(parts):
                para = cells[j].paragraphs[0] if k == 0 else cells[j].add_paragraph()
                para.paragraph_format.space_after = Pt(1)
                add_inline(para, part, size=font)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_callout(doc, lines: list[str]) -> None:
    text = " ".join(l.strip() for l in lines).strip()
    m = re.match(r"\*\*([A-Za-z ]+?):?\*\*:?", text)
    kind = m.group(1).lower().strip() if m else "note"
    fill = CALLOUT_FILLS.get(kind, CALLOUT_FILLS["note"])
    p = doc.add_paragraph()
    _shade(p._p, fill)
    _left_border(p)
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    add_inline(p, text)


def add_code(doc, lines: list[str]) -> None:
    p = doc.add_paragraph()
    _shade(p._p, "F4F4F4")
    p.paragraph_format.space_after = Pt(6)
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = MONO_FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
        r.font.size = Pt(8)
        if i < len(lines) - 1:
            r.add_break()


def add_list_item(doc, text: str, ordered: bool, level: int, number: str | None = None) -> None:
    """Bullets use Word list styles; numbered items keep the Markdown number literally
    (python-docx shares one numbering sequence across the document, which would make
    every numbered list continue from the previous one)."""
    indent = Cm(0.63 * (level + 1))
    if ordered:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = indent
        p.paragraph_format.first_line_indent = Cm(-0.63)
        p.paragraph_format.space_after = Pt(2)
        _run(p, f"{number or '1'}.\t")
        p.paragraph_format.tab_stops.add_tab_stop(indent)
        add_inline(p, text)
        return
    style = "List Bullet" if level == 0 else f"List Bullet {min(level + 1, 3)}"
    try:
        p = doc.add_paragraph(style=style)
    except KeyError:
        p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = indent
    p.paragraph_format.first_line_indent = Cm(-0.63)
    p.paragraph_format.space_after = Pt(2)
    add_inline(p, text)


# ---------------------------------------------------------------- parser
def parse_meta(lines: list[str]) -> tuple[dict, list[str]]:
    i = 0
    while i < len(lines) and not lines[i].strip():
        i += 1
    if i < len(lines) and lines[i].strip() == "---":
        meta, j = {}, i + 1
        while j < len(lines) and lines[j].strip() != "---":
            if ":" in lines[j]:
                k, v = lines[j].split(":", 1)
                meta[k.strip().lower()] = v.strip().strip('"')
            j += 1
        if j < len(lines):
            return meta, lines[j + 1:]
    return {}, lines


def convert(md_text: str, out_path: Path) -> Path:
    lines = md_text.replace("\r\n", "\n").split("\n")
    meta, lines = parse_meta(lines)

    # Use first H1 as title if metadata lacks one; skip that H1 in the body.
    first_h1_idx = next((k for k, l in enumerate(lines) if l.startswith("# ")), None)
    if first_h1_idx is not None:
        meta.setdefault("title", lines[first_h1_idx][2:].strip())
        lines = lines[:first_h1_idx] + lines[first_h1_idx + 1:]

    doc = setup_document()
    add_cover(doc, meta)
    add_header_footer(doc, meta)
    _set_update_fields_on_open(doc)

    i, para_buf = 0, []

    def flush_para():
        nonlocal para_buf
        if para_buf:
            p = doc.add_paragraph()
            add_inline(p, " ".join(s.strip() for s in para_buf))
            para_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush_para()
            i += 1
            continue

        if stripped.startswith("```"):
            flush_para()
            code, i = [], i + 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip("\n"))
                i += 1
            add_code(doc, code)
            i += 1
            continue

        m = HEADING_RE.match(stripped)
        if m:
            flush_para()
            level = min(len(m.group(1)), 4)
            h = doc.add_heading(level=level)
            add_inline(h, m.group(2))
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            flush_para()
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1]):
            flush_para()
            header = split_row(stripped)
            rows, i = [], i + 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, header, rows)
            continue

        if stripped.startswith(">"):
            flush_para()
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip()[1:].strip())
                i += 1
            add_callout(doc, block)
            continue

        bm, nm = BULLET_RE.match(line), NUMBER_RE.match(line)
        if bm or nm:
            flush_para()
            mm = bm or nm
            indent = len(mm.group(1).replace("\t", "    "))
            ordered = bool(nm and not bm)
            number = re.match(r"\s*(\d+)", line).group(1) if ordered else None
            add_list_item(doc, mm.group(2), ordered=ordered, level=min(indent // 2, 2), number=number)
            i += 1
            continue

        para_buf.append(line)
        i += 1

    flush_para()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def main(argv: list[str]) -> int:
    if len(argv) < 2 or argv[1] in ("-h", "--help"):
        print(__doc__)
        return 0 if len(argv) >= 2 else 1
    src = Path(argv[1])
    if not src.exists():
        sys.stderr.write(f"Input not found: {src}\n")
        return 1
    out = Path(argv[2]) if len(argv) > 2 else src.with_suffix(".docx")
    convert(src.read_text(encoding="utf-8"), out)
    print(f"Wrote {out.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
